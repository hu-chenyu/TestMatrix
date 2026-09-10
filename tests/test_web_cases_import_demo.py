"""
TestMatrix Day21: 用例批量导入API测试（POST /api/cases/import）

测试覆盖:
    1. test_import_yaml_success            上传合法YAML（3条）返回200，统计正确
    2. test_import_excel_success           上传合法Excel（3条）返回200，统计正确
    3. test_import_no_file                 未上传文件返回400，message含"未上传"
    4. test_import_unsupported_format      上传.txt返回400，message含"不支持的文件格式"
    5. test_import_yaml_invalid_data        YAML缺必填name返回400，message含字段名
    6. test_import_excel_invalid_priority   Excel priority=P9返回400，message含priority
    7. test_import_upsert_idempotent       相同文件连续导入两次，第二次全量更新
    8. test_import_verify_in_db            导入后列表接口可查到落库数据
    9. test_import_excel_with_sheet_name   双sheet Excel指定sheet_name导入第二个sheet
    10. test_import_multiple_cases         20条大批量YAML导入无异常

测试基建:
    临时SQLite空库（tmp_path + 前后DatabaseSession.reset()防
    Windows文件锁）+ Flask test client端到端验证，零mock零真实外部依赖;
    上传文件由测试自建在tmp_path，接口内部临时文件由接口自行清理。
"""

from pathlib import Path
from typing import Iterator, Optional

import allure
import pytest
import yaml
from flask.testing import FlaskClient
from openpyxl import Workbook

from src.db.db_session import DatabaseSession
from src.web import create_app

# Excel数据文件表头（与DataDriver标准字段对齐，tags/description可选）
EXCEL_HEADERS = ("case_id", "name", "module", "priority", "tags", "description")


# ===========================================================================
# 造数与夹具
# ===========================================================================
def _build_yaml_cases(count: int = 3, prefix: str = "TM-IMP") -> list[dict]:
    """
    构建N条合法YAML用例数据（内部方法）

    数据特征: case_id按prefix+序号生成保证唯一，priority在P0-P3间
    循环覆盖全枚举，tags为列表（YAML原生支持）。

    参数:
        count (int): 生成条数，默认3
        prefix (str): case_id前缀，默认"TM-IMP"

    返回:
        list[dict]: 合法用例字典列表
    """
    priorities = ("P0", "P1", "P2", "P3")
    return [
        {
            "case_id": f"{prefix}-{index:04d}",
            "name": f"批量导入用例{index}",
            "module": "用户中心",
            "priority": priorities[(index - 1) % len(priorities)],
            "tags": ["smoke"],
            "description": f"Day21导入接口验证用例{index}",
        }
        for index in range(1, count + 1)
    ]


def _build_excel_cases(count: int = 3, prefix: str = "TM-IMP-X") -> list[dict]:
    """
    构建N条合法Excel用例数据（内部方法）

    与YAML造数的差异: tags为逗号分隔字符串（openpyxl单元格
    不支持列表写入，DataDriver支持字符串tags自动规范化为列表）。

    参数:
        count (int): 生成条数，默认3
        prefix (str): case_id前缀，默认"TM-IMP-X"

    返回:
        list[dict]: 合法用例字典列表
    """
    priorities = ("P0", "P1", "P2", "P3")
    return [
        {
            "case_id": f"{prefix}-{index:04d}",
            "name": f"Excel导入用例{index}",
            "module": "订单中心",
            "priority": priorities[(index - 1) % len(priorities)],
            "tags": "regression",
            "description": f"Day21 Excel导入验证{index}",
        }
        for index in range(1, count + 1)
    ]


def _make_yaml_file(
    tmp_path: Path, cases: list[dict], filename: str = "cases.yaml"
) -> Path:
    """
    生成YAML用例数据文件（内部方法）

    用yaml.dump将用例列表写入tmp_path下的YAML文件，
    allow_unicode保留中文可读性，顶层为列表格式。

    参数:
        tmp_path (Path): pytest临时目录
        cases (list[dict]): 用例字典列表
        filename (str): 生成文件名，默认"cases.yaml"

    返回:
        Path: 生成的YAML文件路径
    """
    file_path = tmp_path / filename
    with open(file_path, "w", encoding="utf-8") as file_handle:
        yaml.dump(cases, file_handle, allow_unicode=True, sort_keys=False)
    return file_path


def _make_excel_file(
    tmp_path: Path,
    cases: list[dict],
    filename: str = "cases.xlsx",
    sheet_name: str = "cases",
) -> Path:
    """
    生成Excel用例数据文件（内部方法）

    用openpyxl创建workbook，首行写表头（case_id/name/module/
    priority/tags/description），后续行写数据；tags为列表时自动
    转逗号分隔字符串（单元格不支持列表）。

    参数:
        tmp_path (Path): pytest临时目录
        cases (list[dict]): 用例字典列表
        filename (str): 生成文件名，默认"cases.xlsx"
        sheet_name (str): sheet名称，默认"cases"

    返回:
        Path: 生成的Excel文件路径
    """
    file_path = tmp_path / filename
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(list(EXCEL_HEADERS))
    for case in cases:
        row = []
        for header in EXCEL_HEADERS:
            value = case.get(header)
            if header == "tags" and isinstance(value, list):
                value = ",".join(value)
            row.append(value)
        worksheet.append(row)
    workbook.save(file_path)
    workbook.close()
    return file_path


def _post_import(
    client: FlaskClient,
    file_path: Path,
    query_string: Optional[dict] = None,
):
    """
    上传文件调用导入接口（内部方法）

    以multipart/form-data方式POST文件到/api/cases/import，
    支持附加查询参数（sheet_name/creator）。

    参数:
        client (FlaskClient): Flask测试客户端
        file_path (Path): 待上传的数据文件路径
        query_string (dict | None): 附加查询参数，默认None

    返回:
        flask.testing.TestResponse: 接口响应对象
    """
    with open(file_path, "rb") as file_handle:
        return client.post(
            "/api/cases/import",
            data={"file": (file_handle, file_path.name)},
            content_type="multipart/form-data",
            query_string=query_string,
        )


@pytest.fixture
def import_client(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> Iterator[FlaskClient]:
    """
    用例批量导入API测试客户端fixture（临时SQLite空库）

    将数据库指向pytest临时目录独立库文件并完成建表（不预置种子
    数据，导入接口自己造数），前后重置DatabaseSession引擎单例，
    防止Windows下SQLite文件句柄残留导致tmp_path清理失败。

    参数:
        monkeypatch (pytest.MonkeyPatch): 环境变量补丁工具
        tmp_path (Path): pytest临时目录

    返回:
        Iterator[FlaskClient]: yield测试客户端，前后完成引擎重置
    """
    monkeypatch.setenv("TM_DB_TYPE", "sqlite")
    monkeypatch.setenv(
        "TM_DB_SQLITE_PATH", str(tmp_path / "cases_import_api.db")
    )
    DatabaseSession.reset()
    DatabaseSession.init_db()
    yield create_app("test").test_client()
    DatabaseSession.reset()


# ===========================================================================
# 用例批量导入API测试
# ===========================================================================
@allure.feature("用例管理API")
@allure.story("用例批量导入")
@allure.severity(allure.severity_level.CRITICAL)
@pytest.mark.api
@pytest.mark.regression
class TestCasesImportApi:
    """POST /api/cases/import 用例批量导入测试"""

    # ------------------------------------------------------------------
    # 正常导入
    # ------------------------------------------------------------------
    def test_import_yaml_success(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        YAML导入成功: 上传3条合法YAML用例文件，返回200，
        data统计total=3/inserted=3/updated=0，file_name为原始上传文件名
        """
        file_path = _make_yaml_file(tmp_path, _build_yaml_cases(3))
        response = _post_import(import_client, file_path)
        data = response.get_json()

        assert response.status_code == 200, "导入成功应返回200"
        assert data["code"] == 200, "响应体code应为200"
        assert data["data"]["file_name"] == "cases.yaml", (
            "file_name应为原始上传文件名"
        )
        assert data["data"]["total"] == 3, "加载总数应为3"
        assert data["data"]["inserted"] == 3, "空库首次导入应全部新增"
        assert data["data"]["updated"] == 0, "空库首次导入不应有更新"

    def test_import_excel_success(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        Excel导入成功: 上传3条合法Excel用例文件，返回200，
        data统计total=3/inserted=3/updated=0
        """
        file_path = _make_excel_file(tmp_path, _build_excel_cases(3))
        response = _post_import(import_client, file_path)
        data = response.get_json()

        assert response.status_code == 200, "导入成功应返回200"
        assert data["code"] == 200, "响应体code应为200"
        assert data["data"]["file_name"] == "cases.xlsx", (
            "file_name应为原始上传文件名"
        )
        assert data["data"]["total"] == 3, "加载总数应为3"
        assert data["data"]["inserted"] == 3, "空库首次导入应全部新增"
        assert data["data"]["updated"] == 0, "空库首次导入不应有更新"

    # ------------------------------------------------------------------
    # 请求非法
    # ------------------------------------------------------------------
    def test_import_no_file(self, import_client: FlaskClient) -> None:
        """
        未上传文件: POST不携带任何文件字段，返回400，
        message含"未上传"
        """
        response = import_client.post("/api/cases/import")
        data = response.get_json()

        assert response.status_code == 400, "未上传文件应返回400"
        assert data["code"] == 400, "响应体code应为400"
        assert "未上传" in data["message"], "错误信息应说明未上传文件"

    def test_import_unsupported_format(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        格式不支持: 上传.txt后缀文件，返回400，
        message含"不支持的文件格式"
        """
        txt_path = tmp_path / "cases.txt"
        txt_path.write_text("plain text not supported", encoding="utf-8")
        response = _post_import(import_client, txt_path)
        data = response.get_json()

        assert response.status_code == 400, "不支持的后缀应返回400"
        assert data["code"] == 400, "响应体code应为400"
        assert "不支持的文件格式" in data["message"], (
            "错误信息应说明格式不支持"
        )

    # ------------------------------------------------------------------
    # 数据校验失败（DataDriver校验错误透传）
    # ------------------------------------------------------------------
    def test_import_yaml_invalid_data(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        YAML数据校验失败: 用例缺必填字段name，返回400，
        DataDriver校验错误透传，message含字段名name与定位信息
        """
        cases = _build_yaml_cases(1)
        cases[0].pop("name")
        file_path = _make_yaml_file(tmp_path, cases)
        response = _post_import(import_client, file_path)
        data = response.get_json()

        assert response.status_code == 400, "数据校验失败应返回400"
        assert data["code"] == 400, "响应体code应为400"
        assert "name" in data["message"], "错误信息应包含缺失字段名name"

    def test_import_excel_invalid_priority(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        Excel数据校验失败: priority=P9不在P0-P3合法值内，返回400，
        message含字段名priority与行号定位信息
        """
        cases = [
            {
                "case_id": "TM-IMP-X-P9",
                "name": "非法优先级用例",
                "module": "订单中心",
                "priority": "P9",
            }
        ]
        file_path = _make_excel_file(tmp_path, cases)
        response = _post_import(import_client, file_path)
        data = response.get_json()

        assert response.status_code == 400, "数据校验失败应返回400"
        assert data["code"] == 400, "响应体code应为400"
        assert "priority" in data["message"], (
            "错误信息应包含字段名priority"
        )

    # ------------------------------------------------------------------
    # 幂等与落库校验
    # ------------------------------------------------------------------
    def test_import_upsert_idempotent(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        幂等upsert: 同一YAML文件连续导入两次，
        第一次inserted=3/updated=0，第二次inserted=0/updated=3
        """
        file_path = _make_yaml_file(
            tmp_path, _build_yaml_cases(3, prefix="TM-IMP-IDEM")
        )

        first_response = _post_import(import_client, file_path)
        first_data = first_response.get_json()["data"]
        assert first_response.status_code == 200
        assert first_data["inserted"] == 3, "首次导入应全部新增"
        assert first_data["updated"] == 0, "首次导入不应有更新"

        second_response = _post_import(import_client, file_path)
        second_data = second_response.get_json()["data"]
        assert second_response.status_code == 200
        assert second_data["total"] == 3, "第二次加载总数不变"
        assert second_data["inserted"] == 0, "重复导入不应再新增"
        assert second_data["updated"] == 3, "重复导入应全部走更新"

    def test_import_verify_in_db(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        落库校验: 导入3条后调用GET /api/cases/?status=all列表接口，
        total为3且case_id集合与导入数据一致（端到端数据真实落库）
        """
        file_path = _make_yaml_file(
            tmp_path, _build_yaml_cases(3, prefix="TM-IMP-DB")
        )
        import_response = _post_import(import_client, file_path)
        assert import_response.status_code == 200, "导入应成功"

        list_response = import_client.get("/api/cases/?status=all")
        list_data = list_response.get_json()["data"]

        assert list_response.status_code == 200
        assert list_data["total"] == 3, "导入的3条用例应已落库可查"
        assert {item["case_id"] for item in list_data["items"]} == {
            "TM-IMP-DB-0001",
            "TM-IMP-DB-0002",
            "TM-IMP-DB-0003",
        }, "落库用例编号应与导入数据完全一致"

    # ------------------------------------------------------------------
    # 高级场景
    # ------------------------------------------------------------------
    def test_import_excel_with_sheet_name(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        指定sheet导入: 上传含first/second两个sheet的Excel，
        指定sheet_name=second导入第二个sheet的2条数据
        （而非默认活动sheet的1条）
        """
        file_path = tmp_path / "multi_sheet_cases.xlsx"
        workbook = Workbook()

        first_sheet = workbook.active
        first_sheet.title = "first"
        first_sheet.append(list(EXCEL_HEADERS))
        first_sheet.append(
            ["TM-IMP-S1-0001", "第一个sheet用例", "用户中心", "P1", None, "不应被导入"]
        )

        second_sheet = workbook.create_sheet("second")
        second_sheet.append(list(EXCEL_HEADERS))
        second_sheet.append(
            ["TM-IMP-S2-0001", "第二个sheet用例1", "订单中心", "P1", None, None]
        )
        second_sheet.append(
            ["TM-IMP-S2-0002", "第二个sheet用例2", "订单中心", "P2", None, None]
        )

        workbook.save(file_path)
        workbook.close()

        response = _post_import(
            import_client, file_path, query_string={"sheet_name": "second"}
        )
        data = response.get_json()

        assert response.status_code == 200
        assert data["data"]["total"] == 2, "应只导入second sheet的2条数据"
        assert data["data"]["inserted"] == 2, "second sheet数据应全部新增"

        # 落库校验: 库中只有second sheet的数据，无first sheet数据
        list_response = import_client.get("/api/cases/?status=all")
        list_data = list_response.get_json()["data"]
        assert list_data["total"] == 2, "库中应仅有指定sheet的2条用例"
        assert {item["case_id"] for item in list_data["items"]} == {
            "TM-IMP-S2-0001",
            "TM-IMP-S2-0002",
        }, "落库数据应来自second sheet"

    def test_import_multiple_cases(
        self, import_client: FlaskClient, tmp_path: Path
    ) -> None:
        """
        大批量导入: 上传20条用例的YAML文件，返回total=20/inserted=20，
        列表接口total=20，验证大批量导入链路无异常
        """
        file_path = _make_yaml_file(
            tmp_path, _build_yaml_cases(20, prefix="TM-IMP-BULK")
        )
        response = _post_import(import_client, file_path)
        data = response.get_json()

        assert response.status_code == 200, "大批量导入应返回200"
        assert data["code"] == 200, "响应体code应为200"
        assert data["data"]["total"] == 20, "加载总数应为20"
        assert data["data"]["inserted"] == 20, "应全部新增"
        assert data["data"]["updated"] == 0, "空库导入不应有更新"

        # 落库校验: 列表接口按全量状态查总数
        list_response = import_client.get("/api/cases/?status=all&page_size=50")
        list_data = list_response.get_json()["data"]
        assert list_data["total"] == 20, "落库总数应为20"
        assert len(list_data["items"]) == 20, "列表应返回全部20条"
